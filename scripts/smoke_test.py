# -*- coding: utf-8 -*-
"""端到端冒烟测试:导入四类知识库 -> 双检索验证 -> RAG 问答验证 -> 混合检索对比。

用法: python scripts/smoke_test.py   (需先设置环境变量 DASHSCOPE_API_KEY)
验证通过后即可用 streamlit 启动网页端。此文件可删除,不影响项目本身。
"""
import sys
from io import BytesIO
from pathlib import Path

# 项目模块位于 src/ 下，这里把 src/ 加入 sys.path；示例素材位于 scripts/assets/
ROOT = Path(__file__).resolve().parent.parent
SRC_ROOT = ROOT / "src"
sys.path.insert(0, str(SRC_ROOT))
ASSET_DIR = ROOT / "scripts" / "assets"

# Windows 控制台默认 GBK，模型回答里若含 emoji 会触发 UnicodeEncodeError
try:
    sys.stdout.reconfigure(errors="replace")
except (AttributeError, OSError):
    pass

from kbrag.compare_service import MODE_VECTOR_VS_HYBRID, run_compare
from kbrag.document_parser import get_file_type, parse_document
from kbrag.history import get_history
from kbrag.knowledge_base import KnowledgeBaseService
from kbrag.rag_service import RagService
from kbrag.retrievers import get_hybrid_retriever, get_vector_retriever

ASSETS = ["尺码推荐.txt", "洗涤养护.txt", "颜色推荐.txt", "新品推荐.md"]
TEST_SESSION_ID = "smoke_test"


def _build_docx_bytes() -> bytes:
    """用 python-docx 现场生成一个 docx，验证 docx 解析与切片。"""
    from docx import Document

    document = Document()
    document.add_heading("机器人保养周期", level=1)
    document.add_paragraph("建议每周清理一次滚刷，每月检查一次滤网。")
    document.add_paragraph("滤网累计使用 150 小时后建议更换，耗材型号为 X1-F。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pdf_bytes() -> bytes:
    """手工拼一个最小可用的纯文本 PDF，验证 pdf 解析与切片。

    说明:手写 PDF 嵌入中文字体非常复杂,这里用英文内容验证解析链路;
    真实中文 PDF 请在网页端上传验证。
    """
    content = "Size Guide: height 180-190 cm, weight 75-85 kg, choose size XL."
    stream = f"BT /F1 12 Tf 72 720 Td ({content}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += f"{index} 0 obj\n".encode("latin-1") + obj + b"\nendobj\n"

    xref_start = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode("latin-1")
    pdf += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        pdf += f"{offset:010d} 00000 n \n".encode("latin-1")
    pdf += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_start}\n%%EOF\n"
    ).encode("latin-1")
    return bytes(pdf)


def _ingest(service: KnowledgeBaseService, name: str, content: bytes):
    """解析 -> 按类型切片 -> 入库，并打印结果。"""
    file_type = get_file_type(name)
    text = parse_document(content, name)
    result = service.upload_by_str(text, name, file_type=file_type)
    print(f"  {name}({file_type}): {len(text)} 字符 -> {result}")


def step1_ingest():
    print("[1/5] 导入四类示例知识库(会逐块调用嵌入 API)...")
    service = KnowledgeBaseService()

    # txt / md：读取 assets 目录下的静态资产
    for name in ASSETS:
        _ingest(service, name, (ASSET_DIR / name).read_bytes())

    # docx / pdf：脚本现场生成，验证「文件字节 -> 解析 -> 差异化切片 -> 入库」全链路
    _ingest(service, "保养周期.docx", _build_docx_bytes())
    _ingest(service, "size_guide.pdf", _build_pdf_bytes())


def step2_retrieve():
    print("[2/5] 双检索对比: 纯向量 vs 向量+关键词兜底")
    query = "我的身高180，体重165斤，穿什么尺码"

    vector_docs = get_vector_retriever().invoke(query)
    hybrid_docs = get_hybrid_retriever().invoke(query)

    for label, docs in (("纯向量", vector_docs), ("混合检索", hybrid_docs)):
        print(f"  [{label}] 命中 {len(docs)} 条")
        for i, doc in enumerate(docs, 1):
            preview = doc.page_content.strip().replace("\n", " ")[:60]
            print(f"    {i}. {preview} | 类型: {doc.metadata.get('file_type', '未知')}")

    assert vector_docs, "纯向量检索没有命中任何结果"
    assert hybrid_docs, "混合检索没有命中任何结果"


def step3_rag():
    print("[3/5] RAG 问答验证(检索 -> Prompt -> qwen-plus)...")
    # 冒烟测试使用独立会话，避免污染真实用户历史；每次从空历史开始
    get_history(TEST_SESSION_ID).clear()
    session_config = {"configurable": {"session_id": TEST_SESSION_ID}}
    rag = RagService()
    res = rag.chain.invoke(
        {"input": "我的身高180，体重165斤，穿什么尺码？"},
        session_config,
    )
    print("  回答:", res)


def step4_compare():
    print("[4/5] 混合检索对比记录(纯向量 vs 向量+关键词)...")
    record = run_compare(
        "根据文档，身高180体重165斤推荐什么尺码？",
        mode=MODE_VECTOR_VS_HYBRID,
    )
    print(
        f"  纯向量命中 {len(record['sources_vector'])} 条；"
        f"混合检索命中 {len(record['sources_hybrid'])} 条"
    )
    print("  纯向量回答:", record["vector"][:60], "...")
    print("  混合检索回答:", record["hybrid"][:60], "...")


if __name__ == "__main__":
    step1_ingest()
    step2_retrieve()
    step3_rag()
    step4_compare()
    print("[5/5] 全部通过")
